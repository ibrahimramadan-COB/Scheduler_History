"""
Local report engine: queries Snowflake's SCHEDULER_HISTORY view,
reshapes GROUPING SETS results into the same pivot structure the old
Apps Script version used, builds an Excel attachment with openpyxl,
and sends email via SMTP. Reuses the exact same Snowflake connection
logic already proven in scripts/snowflake_writer.py.
"""

import os
import sys
import json
import tempfile

# scripts/ sits one level up from local_app/
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
from snowflake_writer import get_connection  # noqa: E402

SCHEDULER_HISTORY_TABLE = "EXTERNAL_DATA.PUBLIC.SCHEDULER_HISTORY"
EVENT_TYPES = ["Cancelled", "Checked In", "Checked Out", "Created", "Deleted", "Edited", "No Show"]


def run_query(sql: str) -> list[dict]:
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(sql)
        columns = [c[0] for c in cursor.description]
        rows = cursor.fetchall()
        return [dict(zip(columns, r)) for r in rows]
    finally:
        cursor.close()
        conn.close()


def escape_sql(s) -> str:
    return "'" + str(s).replace("'", "''") + "'"


def get_filter_options() -> dict:
    dept_rows = run_query(f"SELECT DISTINCT DEPARTMENT, FUNCTION_NAME FROM {SCHEDULER_HISTORY_TABLE} ORDER BY 1, 2")
    groups = []
    group2_by_group = {}
    for r in dept_rows:
        g, g2 = r["DEPARTMENT"], r["FUNCTION_NAME"]
        if not g:
            continue
        if g not in groups:
            groups.append(g)
        group2_by_group.setdefault(g, [])
        if g2 and g2 not in group2_by_group[g]:
            group2_by_group[g].append(g2)

    clinic_rows = run_query(f"SELECT DISTINCT CLINIC FROM {SCHEDULER_HISTORY_TABLE} WHERE CLINIC IS NOT NULL ORDER BY 1")
    clinics = [r["CLINIC"] for r in clinic_rows if r["CLINIC"]]

    return {
        "groups": sorted(groups),
        "group2ByGroup": group2_by_group,
        "eventTypes": EVENT_TYPES,
        "clinics": clinics,
    }


def build_where_clause(config: dict, event_type_filter, group_filter, group2_filter) -> str:
    clauses = [f"CREATED_TIMESTAMP BETWEEN {escape_sql(config['startDate'])} AND {escape_sql(config['endDate'])}"]
    if event_type_filter:
        clauses.append(f"EVENT_ACTION IN ({', '.join(escape_sql(e) for e in event_type_filter)})")
    if group_filter:
        clauses.append(f"DEPARTMENT = {escape_sql(group_filter)}")
    if group2_filter:
        clauses.append(f"FUNCTION_NAME = {escape_sql(group2_filter)}")
    clinics = [c for c in (config.get("clinics") or []) if c]
    if clinics:
        clauses.append(f"CLINIC IN ({', '.join(escape_sql(c) for c in clinics)})")
    return " AND ".join(clauses)


def reshape_grouping_sets(rows: list[dict], extra_cols=None) -> list[dict]:
    """Turns GROUPING SETS output (detail + dept subtotal + grand total rows,
    identified by which columns are NULL) into header/data/subtotal/grandtotal
    pivot rows, same shape the email/Excel builders expect."""
    extra_cols = extra_cols or []
    out = []
    current_dept = None

    for r in rows:
        is_grand_total = r["DEPARTMENT"] is None
        if is_grand_total:
            out.append({
                "type": "grandtotal", "label": "GRAND TOTAL",
                "patients": int(r["NUM_PATIENTS"] or 0), "appointments": int(r["NUM_APPOINTMENTS"] or 0),
            })
            continue

        dept_key = f'{r["DEPARTMENT"]} — {r["FUNCTION_NAME"]}'
        is_dept_subtotal = r["CREATOR_USER"] is None

        if dept_key != current_dept:
            out.append({"type": "header", "label": dept_key})
            current_dept = dept_key

        if is_dept_subtotal:
            out.append({
                "type": "subtotal", "label": "Total",
                "patients": int(r["NUM_PATIENTS"] or 0), "appointments": int(r["NUM_APPOINTMENTS"] or 0),
            })
        else:
            row = {
                "type": "data", "label": r["CREATOR_USER"],
                "patients": int(r["NUM_PATIENTS"] or 0), "appointments": int(r["NUM_APPOINTMENTS"] or 0),
            }
            if "CREATE_DATE" in r:
                row["createDate"] = r["CREATE_DATE"]
            if "CLINIC" in r:
                row["clinic"] = r["CLINIC"]
            out.append(row)

    return out


def build_html_email(pivot_rows: list[dict], title: str, start_date: str, end_date: str) -> str:
    html = '<div style="font-family: Arial, sans-serif; color: #333;">'
    html += f'<h2 style="color:#1F4E79; margin-bottom:5px;">{title}</h2>'
    html += f'<p style="color:#666; margin-top:0;">Created {start_date} → {end_date}</p>'
    html += ('<table border="1" cellpadding="8" style="border-collapse:collapse; font-size:13px; '
             'text-align:center; max-width:600px;">')
    html += ('<tr><th style="background-color:#203764; color:white; text-align:left;">CREATOR USER</th>'
             '<th style="background-color:#203764; color:white;"># of Patients</th>'
             '<th style="background-color:#203764; color:white;"># of Appointments</th></tr>')

    for row in pivot_rows:
        if row["type"] == "header":
            html += (f'<tr><td colspan="3" style="background-color:#1F4E79; color:white; '
                     f'font-weight:bold; text-align:left;">{row["label"]}</td></tr>')
        elif row["type"] == "data":
            html += (f'<tr><td style="text-align:left; padding-left:25px;">{row["label"]}</td>'
                      f'<td>{row["patients"]}</td><td>{row["appointments"]}</td></tr>')
        elif row["type"] == "subtotal":
            html += (f'<tr style="background-color:#DDEBF7; font-weight:bold;">'
                      f'<td style="text-align:left; padding-left:25px;">{row["label"]}</td>'
                      f'<td>{row["patients"]}</td><td>{row["appointments"]}</td></tr>')
        elif row["type"] == "grandtotal":
            html += (f'<tr style="background-color:#203764; color:white; font-weight:bold;">'
                      f'<td style="text-align:left;">{row["label"]}</td>'
                      f'<td>{row["patients"]}</td><td>{row["appointments"]}</td></tr>')

    html += "</table></div>"
    return html


def _write_pivot_sheet(ws, pivot_rows: list[dict], include_clinic: bool):
    from openpyxl.styles import Font, PatternFill

    header = (["CREATOR USER", "CREATE DATE", "CLINIC", "# of Patients", "# of Appointments"]
              if include_clinic else
              ["CREATOR USER", "CREATE DATE", "# of Patients", "# of Appointments"])
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")

    for row in pivot_rows:
        if row["type"] == "header":
            blank = [""] * len(header)
            blank[0] = row["label"]
            ws.append(blank)
        elif row["type"] == "data":
            if include_clinic:
                ws.append([row["label"], row.get("createDate", ""), row.get("clinic", ""),
                           row["patients"], row["appointments"]])
            else:
                ws.append([row["label"], row.get("createDate", ""), row["patients"], row["appointments"]])
        else:
            blank = [""] * len(header)
            blank[0] = row["label"]
            blank[-2] = row["patients"]
            blank[-1] = row["appointments"]
            ws.append(blank)

    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2


def build_excel_attachment(per_user_rows, per_clinic_rows, report_label: str) -> str:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Per User"
    _write_pivot_sheet(ws1, per_user_rows, include_clinic=False)

    ws2 = wb.create_sheet("Per Clinic")
    _write_pivot_sheet(ws2, per_clinic_rows, include_clinic=True)

    safe_name = "".join(c if c.isalnum() or c in " -_" else "_" for c in report_label)
    tmp_path = os.path.join(tempfile.gettempdir(), f"{safe_name}.xlsx")
    wb.save(tmp_path)
    return tmp_path


def send_email(subject: str, html_body: str, to_list: list[str], attachment_path: str | None = None):
    import smtplib
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.application import MIMEApplication

    smtp_host = os.environ.get("SMTP_HOST", "smtp.gmail.com")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER")
    smtp_password = os.environ.get("SMTP_PASSWORD")
    sender_name = os.environ.get("EMAIL_SENDER_NAME", "WebPT Scheduler History Automation")

    if not smtp_user or not smtp_password:
        raise RuntimeError("SMTP_USER / SMTP_PASSWORD not set. For Gmail, this needs an App Password "
                            "(myaccount.google.com/apppasswords), not your regular password.")

    msg = MIMEMultipart()
    msg["From"] = f"{sender_name} <{smtp_user}>"
    msg["To"] = ", ".join(to_list)
    msg["Subject"] = subject
    msg.attach(MIMEText(html_body, "html"))

    if attachment_path:
        with open(attachment_path, "rb") as f:
            part = MIMEApplication(f.read(), Name=os.path.basename(attachment_path))
        part["Content-Disposition"] = f'attachment; filename="{os.path.basename(attachment_path)}"'
        msg.attach(part)

    with smtplib.SMTP(smtp_host, smtp_port) as server:
        server.starttls()
        server.login(smtp_user, smtp_password)
        server.sendmail(smtp_user, to_list, msg.as_string())


def get_email_list() -> list[str]:
    """Reads recipient emails from the 'Emails' tab in the main Google Sheet —
    kept there rather than duplicated locally, so there's one source of truth."""
    import gspread
    from google.oauth2.service_account import Credentials

    scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
    file_path = os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE")
    if file_path:
        creds = Credentials.from_service_account_file(file_path, scopes=scopes)
    else:
        raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
        if not raw:
            raise RuntimeError("Set GOOGLE_SERVICE_ACCOUNT_FILE or GOOGLE_SERVICE_ACCOUNT_JSON to read the Emails tab.")
        creds = Credentials.from_service_account_info(json.loads(raw), scopes=scopes)

    client = gspread.authorize(creds)
    sheet_id = os.environ.get("SHEET_ID")
    if not sheet_id:
        raise RuntimeError("SHEET_ID env var is not set.")
    spreadsheet = client.open_by_key(sheet_id)
    ws = spreadsheet.worksheet("Emails")
    values = ws.col_values(1)
    return [v.strip() for v in values if "@" in v]


def run_scheduler_report(config: dict) -> str:
    mode = config.get("mode")
    if mode == "created":
        event_type_filter, group_filter, group2_filter = ["Created"], "CC", ""
        report_label = "Created Appointments — CC Department"
    elif mode == "vfd":
        event_type_filter, group_filter, group2_filter = [], "", "VFD"
        report_label = "VFD Activity — All Event Types"
    else:
        event_type_filter = config.get("eventTypes") or []
        group_filter = config.get("group") or ""
        group2_filter = config.get("group2") or ""
        report_label = "Custom Report"

    where_clause = build_where_clause(config, event_type_filter, group_filter, group2_filter)

    summary_sql = f"""
        SELECT DEPARTMENT, FUNCTION_NAME, CREATOR_USER,
               COUNT(DISTINCT PATIENT) AS NUM_PATIENTS, COUNT(*) AS NUM_APPOINTMENTS
        FROM {SCHEDULER_HISTORY_TABLE}
        WHERE {where_clause}
        GROUP BY GROUPING SETS ((DEPARTMENT, FUNCTION_NAME, CREATOR_USER), (DEPARTMENT, FUNCTION_NAME), ())
        ORDER BY DEPARTMENT NULLS LAST, FUNCTION_NAME NULLS LAST, CREATOR_USER NULLS LAST
    """
    summary_rows = run_query(summary_sql)

    grand_total_row = next((r for r in summary_rows if r["DEPARTMENT"] is None), None)
    if not summary_rows or not grand_total_row or not grand_total_row.get("NUM_APPOINTMENTS"):
        return f"⚠️ No rows matched this filter for {config['startDate']} → {config['endDate']}."

    summary_pivot = reshape_grouping_sets(summary_rows)
    html = build_html_email(summary_pivot, report_label, config["startDate"], config["endDate"])
    total_appointments = int(grand_total_row["NUM_APPOINTMENTS"] or 0)

    email_status = ""
    if config.get("sendEmail", True):
        try:
            email_list = get_email_list()
        except Exception as e:
            email_list = []
            email_status = f" (⚠️ Could not read Emails tab: {e})"

        if email_list:
            subject = f"{report_label} ({config['startDate']} to {config['endDate']})"
            attachment_path = None
            if config.get("attachExcel"):
                per_user_sql = f"""
                    SELECT DEPARTMENT, FUNCTION_NAME, CREATOR_USER,
                           TO_VARCHAR(CREATED_TIMESTAMP, 'MM/DD/YYYY') AS CREATE_DATE,
                           COUNT(DISTINCT PATIENT) AS NUM_PATIENTS, COUNT(*) AS NUM_APPOINTMENTS
                    FROM {SCHEDULER_HISTORY_TABLE}
                    WHERE {where_clause}
                    GROUP BY GROUPING SETS ((DEPARTMENT, FUNCTION_NAME, CREATOR_USER, CREATED_TIMESTAMP), (DEPARTMENT, FUNCTION_NAME), ())
                    ORDER BY DEPARTMENT NULLS LAST, FUNCTION_NAME NULLS LAST, CREATOR_USER NULLS LAST, CREATED_TIMESTAMP NULLS LAST
                """
                per_clinic_sql = f"""
                    SELECT DEPARTMENT, FUNCTION_NAME, CREATOR_USER,
                           TO_VARCHAR(CREATED_TIMESTAMP, 'MM/DD/YYYY') AS CREATE_DATE, CLINIC,
                           COUNT(DISTINCT PATIENT) AS NUM_PATIENTS, COUNT(*) AS NUM_APPOINTMENTS
                    FROM {SCHEDULER_HISTORY_TABLE}
                    WHERE {where_clause}
                    GROUP BY GROUPING SETS ((DEPARTMENT, FUNCTION_NAME, CREATOR_USER, CREATED_TIMESTAMP, CLINIC), (DEPARTMENT, FUNCTION_NAME), ())
                    ORDER BY DEPARTMENT NULLS LAST, FUNCTION_NAME NULLS LAST, CREATOR_USER NULLS LAST, CREATED_TIMESTAMP NULLS LAST, CLINIC NULLS LAST
                """
                per_user_pivot = reshape_grouping_sets(run_query(per_user_sql), ["CREATE_DATE"])
                per_clinic_pivot = reshape_grouping_sets(run_query(per_clinic_sql), ["CREATE_DATE", "CLINIC"])
                attachment_path = build_excel_attachment(per_user_pivot, per_clinic_pivot, report_label)

            send_email(subject, html, email_list, attachment_path)
            email_status = f" Emailed to: {', '.join(email_list)}." + (" (Excel attached)" if attachment_path else "")
        elif not email_status:
            email_status = " (⚠️ No valid addresses found in 'Emails' tab — skipped sending.)"

    return f"✅ {report_label}: {total_appointments} appointment(s) for {config['startDate']} → {config['endDate']}.{email_status}"
